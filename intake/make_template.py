#!/usr/bin/env python3
"""Generate the Excel strategy template (strategy_template.xlsx), pre-filled with a worked example."""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from spec import INDICATOR_TYPES

HEAD = PatternFill("solid", fgColor="1C5CAB")
SOFT = PatternFill("solid", fgColor="EEF3FA")
HFONT = Font(bold=True, color="FFFFFF", name="Arial", size=11)
BFONT = Font(name="Arial", size=11)
MONO = Font(name="Consolas", size=11)
LINE = Border(bottom=Side(style="thin", color="D9DDE3"))


def sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.sheet_view.rightToLeft = True
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.alignment = HEAD, HFONT, Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    return ws


def rows(ws, data, mono_cols=()):
    for r, row in enumerate(data, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = MONO if c in mono_cols else BFONT
            cell.border = LINE
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build(path):
    wb = Workbook()
    ws = sheet(wb, "الاستراتيجية", ["الحقل", "القيمة", "ملاحظة"], [22, 46, 52], first=True)
    rows(ws, [
        ["الاسم", "rsi2_trend_pullback", "بالإنجليزية snake_case: حروف صغيرة وأرقام و _"],
        ["العائلة", "ارتداد", "اتجاه / ارتداد / جلسات / أخرى"],
        ["الوصف", "شراء التصحيحات القصيرة داخل اتجاه صاعد (RSI 2)", "سطر يشرح الفكرة"],
        ["المصدر", "quantpedia", "quantpedia, ssrn, arxiv, quantconnect, tradingview, mql5, github, forexfactory, babypips, other"],
        ["الرابط", "https://quantpedia.com", "رابط المصدر الأصلي إن وُجد"],
        ["الادعاء", "نسبة ربح عالية", "الأداء المعلن كما ورد في المصدر"],
        ["الترخيص", "", "اتركه فارغًا إن لم يكن معروفًا"],
        ["وقف الخسارة (ATR)", 0, "مضاعف ATR لوقف الخسارة، 0 = بدون"],
    ], mono_cols=(2,))
    fam = DataValidation(type="list", formula1='"اتجاه,ارتداد,جلسات,أخرى"', allow_blank=False)
    ws.add_data_validation(fam)
    fam.add("B3")

    ws = sheet(wb, "المؤشرات", ["name", "type", "period", "source", "k", "fast", "slow", "signal"], [16, 14, 12, 12, 8, 8, 8, 8])
    rows(ws, [
        ["r", "rsi", "{rsi_n}", "close", None, None, None, None],
        ["trend", "sma", 200, "close", None, None, None, None],
        ["exit_ma", "sma", 5, "close", None, None, None, None],
    ], mono_cols=(1, 2, 3, 4))
    dv = DataValidation(type="list", formula1='"' + ",".join(INDICATOR_TYPES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2:B40")
    src = DataValidation(type="list", formula1='"open,high,low,close"', allow_blank=True)
    ws.add_data_validation(src)
    src.add("D2:D40")

    ws = sheet(wb, "القواعد", ["القاعدة", "الشرط", "شرح"], [20, 58, 40])
    rows(ws, [
        ["دخول شراء", "r < {band} and close > trend", "RSI منخفض جدًا والسعر فوق متوسط 200"],
        ["خروج من الشراء", "close > exit_ma", "الإغلاق فوق متوسط 5"],
        ["دخول بيع", "r > 100 - {band} and close < trend", "العكس"],
        ["خروج من البيع", "close < exit_ma", ""],
    ], mono_cols=(2,))

    ws = sheet(wb, "المعاملات", ["المعامل", "القيم", "ملاحظة"], [18, 26, 50])
    rows(ws, [
        ["rsi_n", "2", "القيمة الأصلية أولًا"],
        ["band", "5, 10", "حد أقصى 8 توليفات لكل المعاملات معًا"],
    ], mono_cols=(1, 2))

    ws = sheet(wb, "دليل", ["البند", "الشرح"], [26, 90])
    guide = [
        ["الفكرة", "صف استراتيجيتك هنا بقواعد واضحة، ثم ارفع الملف من تبويب «إضافة استراتيجية» في اللوحة. يُحوَّل مباشرة بلا ذكاء اصطناعي."],
        ["توقيت القرار", "كل شرط يُقيَّم عند إغلاق الشمعة، والتنفيذ على الشمعة التالية. لا يمكن النظر إلى المستقبل."],
        ["المتغيرات الجاهزة", "open, high, low, close, hour (UTC), weekday (0=الاثنين), day, month"],
        ["الإزاحة للماضي", "close[1] = إغلاق الشمعة السابقة. الأرقام السالبة مرفوضة."],
        ["المقارنات", "<  <=  >  >=  ==  !="],
        ["المنطق", "and  or  not  والأقواس ( )"],
        ["العمليات", "+  -  *  /   مثل: (close - close[24]) / close[24] > 0.01"],
        ["التقاطع", "crosses_above(fast, slow)   crosses_below(fast, slow)"],
        ["القيمة المطلقة", "abs(x)"],
        ["المعاملات", "اكتب {اسم_المعامل} داخل الشروط أو الفترات، وعرّف قيمه في ورقة «المعاملات»."],
        ["أنواع المؤشرات", ", ".join(INDICATOR_TYPES)],
        ["bb_upper / bb_lower", "تحتاج period و k (عدد الانحرافات المعيارية)"],
        ["macd / macd_signal", "تحتاج fast و slow (و signal للخط الإشاري)"],
        ["roc", "نسبة التغير % خلال period شمعة"],
        ["zscore", "بُعد السعر عن متوسطه بوحدات الانحراف المعياري"],
        ["الخروج", "إن لم تكتب قاعدة خروج، يبقى المركز حتى ظهور إشارة الاتجاه المعاكس."],
        ["ممنوع", "Martingale، Grid، تعزيز المراكز الخاسرة — لا يمكن التعبير عنها هنا عمدًا."],
        ["مثال: تقاطع متوسطات", "مؤشرات: fast=ema 20، slow=ema 50 — دخول شراء: crosses_above(fast, slow) — دخول بيع: crosses_below(fast, slow)"],
        ["مثال: كسر جلسة لندن", "مؤشرات: hi=highest 7 (source=high) — دخول شراء: hour == 7 and close > hi[1] — خروج: hour >= 16"],
    ]
    rows(ws, guide)
    wb.save(path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "strategy_template.xlsx")
    print("template written")
