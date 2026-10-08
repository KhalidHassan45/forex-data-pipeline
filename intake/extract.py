"""
Turn an uploaded file into (a) a spec directly (Excel template), (b) a Python port (lab contract),
or (c) plain text for the translator. Never executes anything from the file.
"""
import io
import re
import zipfile
from pathlib import Path

CODE_EXT = {".pine": "pine", ".mq4": "mql4", ".mq5": "mql5", ".mqh": "mql4", ".py": "python",
            ".txt": "text", ".md": "text", ".json": "json"}
SHEET_EXT = {".xlsx": "excel", ".xlsm": "excel", ".csv": "csv"}
DOC_EXT = {".pdf": "pdf", ".docx": "docx"}
ALLOWED = set(CODE_EXT) | set(SHEET_EXT) | set(DOC_EXT)
MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 60_000          # characters sent to the translator
MAX_PDF_PAGES = 40

RULE_ALIASES = {
    "long_entry": "long_entry", "دخول شراء": "long_entry", "شراء": "long_entry",
    "long_exit": "long_exit", "خروج من الشراء": "long_exit", "خروج شراء": "long_exit",
    "short_entry": "short_entry", "دخول بيع": "short_entry", "بيع": "short_entry",
    "short_exit": "short_exit", "خروج من البيع": "short_exit", "خروج بيع": "short_exit",
}
META_ALIASES = {
    "name": "name", "الاسم": "name", "family": "family", "العائلة": "family",
    "description": "description", "الوصف": "description", "source": "source", "المصدر": "source",
    "url": "url", "الرابط": "url", "claim": "claim", "الادعاء": "claim", "license": "license", "الترخيص": "license",
    "stop_atr": "stop_atr", "وقف الخسارة (ATR)": "stop_atr", "وقف الخسارة": "stop_atr",
}
FAMILY_AR = {"اتجاه": "trend", "ارتداد": "reversion", "جلسات": "session", "أخرى": "other"}


class ExtractError(ValueError):
    pass


def kind_of(filename):
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED:
        raise ExtractError(f"نوع الملف {ext or '(بلا امتداد)'} غير مدعوم. المدعوم: {', '.join(sorted(ALLOWED))}")
    return ext, CODE_EXT.get(ext) or SHEET_EXT.get(ext) or DOC_EXT.get(ext)


def sniff(data: bytes, ext: str):
    """Reject files whose bytes do not match their extension (e.g. an .exe renamed .pdf)."""
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise ExtractError("الملف لا يبدو PDF حقيقيًا")
    if ext in (".xlsx", ".xlsm", ".docx") and not data.startswith(b"PK"):
        raise ExtractError("الملف لا يبدو ملف Office حقيقيًا")
    if ext in CODE_EXT or ext == ".csv":
        if b"\x00" in data[:4096]:
            raise ExtractError("ملف نصي يحتوي بيانات ثنائية — رُفض")


def text_of_code(data: bytes):
    for enc in ("utf-8", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def pdf_text(data: bytes):
    from pypdf import PdfReader
    try:
        r = PdfReader(io.BytesIO(data))
        if r.is_encrypted:
            raise ExtractError("ملف PDF محمي بكلمة مرور")
        pages = r.pages[:MAX_PDF_PAGES]
        txt = "\n\n".join(f"[صفحة {i + 1}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(pages))
    except ExtractError:
        raise
    except Exception as e:
        raise ExtractError(f"تعذّرت قراءة PDF: {e}")
    if len(re.sub(r"\[صفحة \d+\]|\s", "", txt)) < 200:
        raise ExtractError("لا يحتوي PDF نصًا قابلًا للقراءة (غالبًا ممسوح ضوئيًا). ارفع نسخة نصية أو صف القواعد في قالب Excel.")
    return txt, len(r.pages)


def docx_text(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
    except Exception as e:
        raise ExtractError(f"تعذّرت قراءة Word: {e}")
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)


def _cell(v):
    return "" if v is None else str(v).strip()


def excel_workbook(data: bytes):
    from openpyxl import load_workbook
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ExtractError(f"تعذّرت قراءة Excel: {e}")


def _sheet(wb, *names):
    for n in wb.sheetnames:
        if n.strip().lower() in names:
            return wb[n]
    return None


def is_template(wb):
    return _sheet(wb, "القواعد", "rules") is not None and _sheet(wb, "الاستراتيجية", "strategy") is not None


def excel_to_spec(wb):
    """Parse the official template (see make_template.py) into (spec, meta)."""
    meta, spec = {}, {"indicators": {}, "rules": {}, "params": {}}
    for row in _sheet(wb, "الاستراتيجية", "strategy").iter_rows(min_row=1, values_only=True):
        k, v = _cell(row[0] if row else ""), _cell(row[1] if row and len(row) > 1 else "")
        key = META_ALIASES.get(k) or META_ALIASES.get(k.lower())
        if key and v:
            meta[key] = v
    ind_sheet = _sheet(wb, "المؤشرات", "indicators")
    if ind_sheet is not None:
        rows = list(ind_sheet.iter_rows(values_only=True))
        head = [_cell(h).lower() for h in rows[0]] if rows else []
        for r in rows[1:]:
            rec = {head[i]: _cell(v) for i, v in enumerate(r) if i < len(head) and _cell(v) != ""}
            name = rec.pop("name", None) or rec.pop("الاسم", None)
            if not name:
                continue
            ind = {}
            for k, v in rec.items():
                k = {"النوع": "type", "الفترة": "period", "المصدر": "source"}.get(k, k)
                ind[k] = v
            spec["indicators"][name.lower()] = ind
    for r in _sheet(wb, "القواعد", "rules").iter_rows(min_row=1, values_only=True):
        k, v = _cell(r[0] if r else ""), _cell(r[1] if r and len(r) > 1 else "")
        key = RULE_ALIASES.get(k) or RULE_ALIASES.get(k.lower())
        if key and v:
            spec["rules"][key] = v
    p_sheet = _sheet(wb, "المعاملات", "params")
    if p_sheet is not None:
        for r in list(p_sheet.iter_rows(values_only=True))[1:]:
            k, v = _cell(r[0] if r else ""), _cell(r[1] if r and len(r) > 1 else "")
            if k and v:
                vals = []
                for x in re.split(r"[,،;\s]+", v):
                    if x:
                        try:
                            f = float(x)
                            vals.append(int(f) if f == int(f) else f)
                        except ValueError:
                            raise ExtractError(f"قيمة غير رقمية «{x}» للمعامل {k}")
                spec["params"][k.lower()] = vals
    name = re.sub(r"[^a-z0-9_]+", "_", meta.get("name", "excel_strategy").lower()).strip("_") or "excel_strategy"
    if not name[0].isalpha():
        name = "s_" + name
    spec["name"] = name[:40]
    spec["family"] = FAMILY_AR.get(meta.get("family", ""), meta.get("family", "other").lower())
    spec["description"] = meta.get("description", "")
    if meta.get("stop_atr"):
        try:
            spec["stop_atr"] = float(meta["stop_atr"])
        except ValueError:
            raise ExtractError("قيمة وقف الخسارة (ATR) يجب أن تكون رقمًا")
    return spec, meta


def sheet_text(wb, max_rows=150):
    parts = []
    for ws in wb.worksheets[:6]:
        parts.append(f"## ورقة: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append("…")
                break
            cells = [_cell(c) for c in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def csv_text(data: bytes, max_lines=200):
    lines = text_of_code(data).splitlines()
    return "\n".join(lines[:max_lines]) + ("\n…" if len(lines) > max_lines else "")


def is_lab_python(src: str):
    return bool(re.search(r"^STRATEGY\s*=", src, re.M))


def extract(filename: str, data: bytes):
    """
    → dict(kind, lang, route, text?, spec?, meta?, pages?)
    route: "template" (spec ready) | "python_port" | "translate"
    """
    if len(data) > MAX_BYTES:
        raise ExtractError("حجم الملف أكبر من 10MB")
    if not data:
        raise ExtractError("الملف فارغ")
    ext, kind = kind_of(filename)
    sniff(data, ext)
    out = {"ext": ext, "kind": kind}
    if kind == "excel":
        wb = excel_workbook(data)
        if is_template(wb):
            spec, meta = excel_to_spec(wb)
            return out | {"route": "template", "spec": spec, "meta": meta, "lang": "excel-template"}
        return out | {"route": "translate", "text": sheet_text(wb)[:MAX_TEXT], "lang": "excel"}
    if kind == "csv":
        return out | {"route": "translate", "text": csv_text(data)[:MAX_TEXT], "lang": "csv"}
    if kind == "pdf":
        txt, pages = pdf_text(data)
        return out | {"route": "translate", "text": txt[:MAX_TEXT], "pages": pages, "lang": "paper",
                      "truncated": len(txt) > MAX_TEXT}
    if kind == "docx":
        txt = docx_text(data)
        return out | {"route": "translate", "text": txt[:MAX_TEXT], "lang": "text"}
    src = text_of_code(data)
    if kind == "python" and is_lab_python(src):
        return out | {"route": "python_port", "text": src, "lang": "python"}
    return out | {"route": "translate", "text": src[:MAX_TEXT], "lang": kind, "truncated": len(src) > MAX_TEXT}
